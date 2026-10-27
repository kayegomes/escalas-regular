import json
import os
import tempfile
import unittest
from pathlib import Path

import pandas as pd

from app_support import (
    OutputFileLockedError,
    append_execution_history,
    fallback_path,
    is_valid_email,
    safe_filename,
    save_with_fallback,
    split_email_addresses,
)
from engine_2405 import _event_score, _normalize_channel, _normalize_text
from engine_2468 import _is_ge_tv_row, _team_key, process_2468_base
from engine_2468 import _is_quickhold_in_scale as _is_quickhold_in_scale_2468
from engine_cross import _score_grade_match
from engine_grades import _consolidate_grade_dataframe_windows, _extend_grade_windows_to_next_event, _merge_repeated_grade_windows, _time_key, extract_sportv_channel_block, flatten_sportv_grade, process_premiere_grade
try:
    from gerador_escalas_desktop import (
        GeradorEscalasApp,
        _build_elenco_value,
        _build_email_subject,
        _build_group_email_subject,
        _combine_html_bodies,
        _extract_html_period,
        _group_html_files_by_recipient,
        _html_delivery_metadata,
        _is_empty_transition_row,
        _split_scale_weeks,
    )
except ModuleNotFoundError:
    GeradorEscalasApp = None
    _build_elenco_value = None
    _build_email_subject = None
    _build_group_email_subject = None
    _combine_html_bodies = None
    _extract_html_period = None
    _group_html_files_by_recipient = None
    _html_delivery_metadata = None
    _is_empty_transition_row = None
    _split_scale_weeks = None
from engine_cross import (
    _append_pre_review_alerts,
    _mark_missing_grade,
    _find_best_grade_match,
    _is_folga_row,
    _is_programa,
    _is_quickhold_in_scale,
    _is_time_to_confirmar,
    _is_valid_time_str,
)


class AppSupportTests(unittest.TestCase):
    def test_safe_filename_removes_invalid_windows_characters(self):
        self.assertEqual(safe_filename('Ana: "Teste"/2026'), "Ana_Teste_2026")
        self.assertEqual(safe_filename("   "), "sem_nome")

    def test_email_validation(self):
        self.assertTrue(is_valid_email("pessoa@example.com"))
        self.assertFalse(is_valid_email("sem-email"))
        self.assertFalse(is_valid_email(""))
        self.assertEqual(
            split_email_addresses("a@example.com; b@example.com, a@example.com; inválido"),
            ["a@example.com", "b@example.com"],
        )

    def test_save_with_fallback_when_original_is_locked(self):
        calls = []

        def writer(path):
            calls.append(path)
            if len(calls) == 1:
                raise PermissionError("locked")
            with open(path, "w", encoding="utf-8") as file:
                file.write("ok")

        with tempfile.TemporaryDirectory() as directory:
            original = os.path.join(directory, "saida.xlsx")
            path, used_fallback = save_with_fallback(writer, original)
            self.assertEqual(path, fallback_path(original))
            self.assertTrue(used_fallback)
            self.assertTrue(os.path.exists(path))

    def test_save_with_fallback_raises_specific_error_when_both_are_locked(self):
        def writer(_path):
            raise PermissionError("locked")

        with tempfile.TemporaryDirectory() as directory:
            original = os.path.join(directory, "saida.xlsx")
            with self.assertRaises(OutputFileLockedError):
                save_with_fallback(writer, original)

    def test_history_is_json_lines(self):
        with tempfile.TemporaryDirectory() as directory:
            history = append_execution_history(directory, "Teste", "SUCESSO", "1 item")
            self.assertIsNotNone(history)
            with open(history, encoding="utf-8") as file:
                entry = json.loads(file.readline())
            self.assertEqual(entry["etapa"], "Teste")
            self.assertEqual(entry["status"], "SUCESSO")


class EngineRegressionTests(unittest.TestCase):
    def test_channel_and_text_normalization(self):
        self.assertEqual(_normalize_channel("Sportv 1"), "SPORTV")
        self.assertEqual(_normalize_channel("TV Globo Rede"), "TV GLOBO")
        self.assertEqual(_normalize_text("Irã x Alemanha"), "IRA X ALEMANHA")

    @unittest.skipIf(_is_empty_transition_row is None, "Tkinter não disponível neste ambiente")
    def test_empty_transition_rows_are_filtered_without_removing_real_activities(self):
        empty = pd.Series({
            "Data": "10/08/2026 para 11/08/2026",
            "Plataforma": "-",
            "Evento/Programa": "-",
            "Produto (WO/Quick Hold)": "-",
            "Local": "-",
            "Início": "00:00",
            "Fim": "00:00",
        })
        real = pd.Series({
            "Data": "15/08/2026",
            "Plataforma": "Sportv",
            "Evento/Programa": "PANELA",
            "Produto (WO/Quick Hold)": "PANELA SPORTV/NA/NA",
            "Local": "Berrini",
            "Início": "23:30",
            "Fim": "01:30",
        })
        self.assertTrue(_is_empty_transition_row(empty))
        self.assertFalse(_is_empty_transition_row(real))

    @unittest.skipIf(_build_elenco_value is None, "Tkinter não disponível neste ambiente")
    def test_team_key_separates_event_windows(self):
        base = {
            "WO#": "WO-1",
            "Data": "16/08/2026",
            "Plataforma": "Sportv 3",
            "Evento/Programa": "CIRCUITO MUNDIAL DE SURF - ETAPA DE TEAHUPOO",
        }
        first = pd.Series({**base, "Início": "14:00", "Fim": "19:00"})
        second = pd.Series({**base, "Início": "19:00", "Fim": "00:00"})
        self.assertNotEqual(_team_key(first), _team_key(second))

    def test_team_grouping_marks_remote_by_narrator(self):
        # "Local de Locução" = "Narração By SP/RJ" identifica o narrador ou
        # comentarista reserva que atua remotamente, em vez do titular
        # ("Local"). Sem essa marca fica impossível diferenciar titular de
        # by no agrupamento concatenado da equipe do evento.
        rows = pd.DataFrame([
            {
                "Nome": "Everaldo Marques", "Tipo de Atividade": "Booking", "WO#": "2644679-1",
                "Data": "18/10/2026", "Início": "15:40", "Fim": "18:10", "Plataforma": "TV Globo",
                "Evento/Programa": "FLUMINENSE X SANTOS", "Equipe": "Esp - Talentos",
                "Função": "Narrador", "Local de Locução": "Local",
            },
            {
                "Nome": "Isabelly Morais", "Tipo de Atividade": "Booking", "WO#": "2644679-1",
                "Data": "18/10/2026", "Início": "15:40", "Fim": "18:10", "Plataforma": "TV Globo",
                "Evento/Programa": "FLUMINENSE X SANTOS", "Equipe": "Esp - Talentos",
                "Função": "Narrador", "Local de Locução": "Narração By SP",
            },
        ])
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "2468.xlsx")
            rows.to_excel(path, index=False)
            df = process_2468_base(path)
        narrador = df.loc[df["Nome"] == "Everaldo Marques", "Narrador"].iloc[0]
        self.assertEqual(narrador, "Everaldo Marques ; Isabelly Morais (by)")

    @unittest.skipIf(_build_elenco_value is None, "Tkinter não disponível neste ambiente")
    def test_elenco_excludes_the_recipient(self):
        row = pd.Series({
            "Elenco": "Grafite ; Rafaelle Seraphim",
            "Narrador": "Grafite",
            "Comentarista": "Rafaelle Seraphim",
            "Repórter": "",
        })
        self.assertEqual(_build_elenco_value(row, "Grafite"), "Rafaelle Seraphim")

    @unittest.skipIf(GeradorEscalasApp is None, "Tkinter não disponível neste ambiente")
    def test_html_keeps_folga_rows_and_uses_description(self):
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp:
            df = pd.DataFrame([{
                "Nome": "Abel Neto",
                "Data": "13/08/2026",
                "Dia": "Thursday",
                "Descrição": "Day Off / Folga",
                "Row Display": "Day Off / Folga",
                "Tipo de Atividade": "Other Time Off",
                "Plataforma": "-",
                "Início": "00:00",
                "Fim": "00:00",
                "Pré": "-",
                "Elenco": "-",
                "Narrador": "-",
                "Comentarista": "Abel Neto",
                "Repórter": "-",
                "Coordenador": "-",
                "Produtor": "-",
                "Local": "Estádio Maracanã - Jornalista Mário Filho",
                "Local de Locução": "Narração By JB",
                "Produto (WO/Quick Hold)": "-",
            }])
            GeradorEscalasApp.gerar_html(object(), "Abel Neto", df, tmp)
            html = next(Path(tmp).glob("escala_*.html")).read_text(encoding="utf-8")
            self.assertIn(">FOLGA<", html)
            self.assertNotIn("Day Off / Folga", html)
            self.assertNotIn(">00:00<", html)

    @unittest.skipIf(GeradorEscalasApp is None, "Tkinter não disponível neste ambiente")
    def test_two_week_scale_generates_separate_preview_html(self):
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp:
            dates = pd.date_range("2026-08-03", periods=14, freq="D")
            df = pd.DataFrame({
                "Nome": ["André Felipe"] * 14,
                "Data": dates.strftime("%d/%m/%Y"),
                "Data_obj": dates,
                "Data_fim_obj": dates,
                "Dia": dates.strftime("%A"),
                "Plataforma": ["Sportv"] * 14,
                "Pré": ["-"] * 14,
                "Início": ["12:00"] * 14,
                "Fim": ["13:00"] * 14,
                "Evento/Programa": [f"Evento {i}" for i in range(14)],
                "Produto (WO/Quick Hold)": ["Produto"] * 14,
                "Local de Locução": ["Narração By JB"] * 14,
                "Elenco": ["Vander Carioca"] * 14,
                "Narrador": ["André Felipe"] * 14,
                "Comentarista": ["Vander Carioca"] * 14,
                "Repórter": ["-"] * 14,
                "Coordenador": ["-"] * 14,
                "Produtor": ["-"] * 14,
            })
            weeks = _split_scale_weeks(df)
            self.assertEqual(len(weeks), 2)
            self.assertEqual(len(weeks[0][1]), 7)
            self.assertEqual(len(weeks[1][1]), 7)
            self.assertEqual(weeks[0][2].strftime("%d/%m/%Y"), "03/08/2026")
            self.assertEqual(weeks[0][3].strftime("%d/%m/%Y"), "09/08/2026")
            self.assertEqual(weeks[1][2].strftime("%d/%m/%Y"), "10/08/2026")
            self.assertEqual(weeks[1][3].strftime("%d/%m/%Y"), "16/08/2026")
            self.assertEqual(len(_split_scale_weeks(df.iloc[:7].copy())), 1)

            GeradorEscalasApp.gerar_html(
                object(), "André Felipe", weeks[0][1], tmp,
                week_number=1, is_preview=False,
                period_start=weeks[0][2], period_end=weeks[0][3],
            )
            GeradorEscalasApp.gerar_html(
                object(), "André Felipe", weeks[1][1], tmp,
                week_number=2, is_preview=True,
                period_start=weeks[1][2], period_end=weeks[1][3],
            )
            html_files = sorted(Path(tmp).glob("escala_*.html"))
            self.assertEqual(len(html_files), 2)
            preview_path = next(path for path in html_files if "previa_semana_2" in path.name)
            preview_html = preview_path.read_text(encoding="utf-8")
            self.assertIn("Prévia da sua escala", preview_html)
            self.assertIn("Prévia da sua escala 10/08/2026 a 16/08/2026", preview_html)
            self.assertNotIn("Oi André Felipe", preview_html)
            self.assertNotIn("Dúvidas ou problemas? É só nos procurar:", preview_html)
            self.assertNotIn("ATENÇÃO:", preview_html)
            self.assertEqual(_extract_html_period(preview_html), "10/08/2026 a 16/08/2026")
            self.assertEqual(
                _html_delivery_metadata(preview_path.name),
                ("André Felipe", 2, True),
            )
            self.assertEqual(
                _build_email_subject("André Felipe", False, "03/08/2026 a 09/08/2026"),
                "Escala - André Felipe (03/08/2026 a 09/08/2026)",
            )
            self.assertEqual(
                _build_email_subject("André Felipe", True, "10/08/2026 a 16/08/2026"),
                "Prévia da sua escala - André Felipe (10/08/2026 a 16/08/2026)",
            )

            grouped = _group_html_files_by_recipient(tmp)
            self.assertEqual(list(grouped), ["André Felipe"])
            deliveries = grouped["André Felipe"]
            self.assertEqual(len(deliveries), 2)
            bodies = []
            for _, _, filename, _ in deliveries:
                bodies.append((Path(tmp) / filename).read_text(encoding="utf-8"))
            combined = _combine_html_bodies(bodies)
            self.assertEqual(combined.lower().count("escala consolidada 03/08/2026 a 09/08/2026"), 1)
            self.assertEqual(combined.lower().count("prévia da sua escala 10/08/2026 a 16/08/2026"), 1)
            self.assertEqual(combined.count("<section class=\"escala-semana\">"), 2)
            self.assertEqual(combined.count("Oi André Felipe"), 1)
            self.assertEqual(combined.count("Dúvidas ou problemas? É só nos procurar:"), 1)
            self.assertEqual(combined.count("ATENÇÃO:"), 1)
            self.assertEqual(combined.count("<html>"), 1)
            self.assertIn("Escala e prévia - André Felipe", _build_group_email_subject("André Felipe", deliveries))

            import logging
            import gerador_escalas_desktop as stage3_module

            class DummyMail:
                def __init__(self):
                    self.Subject = ""
                    self.To = ""
                    self.HTMLBody = ""
                    self.CC = ""
                    self.displayed = False

                def Display(self):
                    self.displayed = True

            class DummyOutlook:
                def __init__(self):
                    self.mails = []

                def CreateItem(self, _item_type):
                    mail = DummyMail()
                    self.mails.append(mail)
                    return mail

            class DummyWin32:
                def __init__(self):
                    self.outlook = DummyOutlook()

                def Dispatch(self, _name):
                    return self.outlook

            class DummyApp:
                def __init__(self):
                    self.logger = logging.getLogger("test_stage3_grouped_email")
                    self.logs = []
                    self.progress = []
                    self.config = {"email": {"cc": "#GESTAO-TRANSMISSOES-PROGRAMAS"}}

                def log(self, message, *args, **kwargs):
                    self.logs.append(str(message))

                def show_error(self, *args, **kwargs):
                    raise AssertionError(f"show_error inesperado: {args}")

                def set_progress(self, value, message):
                    self.progress.append((value, message))

            previous_win32 = stage3_module.win32_client
            fake_win32 = DummyWin32()
            stage3_module.win32_client = fake_win32
            try:
                result = GeradorEscalasApp.enviar_emails(
                    DummyApp(),
                    tmp,
                    contacts={"andré felipe": "andre@example.com; second@example.com; andre@example.com"},
                )
            finally:
                stage3_module.win32_client = previous_win32
            self.assertEqual(result["encontrados"], 2)
            self.assertEqual(result["rascunhos"], 1)
            self.assertEqual(result["ignorados"], 0)
            self.assertEqual(len(fake_win32.outlook.mails), 1)
            mail = fake_win32.outlook.mails[0]
            self.assertTrue(mail.displayed)
            self.assertEqual(mail.To, "andre@example.com; second@example.com")
            self.assertEqual(mail.CC, "#GESTAO-TRANSMISSOES-PROGRAMAS")
            self.assertIn("Escala e prévia - André Felipe", mail.Subject)
            self.assertEqual(mail.HTMLBody.lower().count("escala consolidada 03/08/2026 a 09/08/2026"), 1)
            self.assertEqual(mail.HTMLBody.lower().count("prévia da sua escala 10/08/2026 a 16/08/2026"), 1)

    def test_html_preserves_elenco_product_and_event_fields(self):
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp:
            df = pd.DataFrame([{
                "Nome": "André Felipe",
                "Plataforma": "Sportv 2",
                "Data": "10/08/2026",
                "Dia": "Monday",
                "Pré": "-",
                "Início": "12:00",
                "Fim": "13:45",
                "Evento/Programa": "EQUADOR X BRASIL",
                "Produto (WO/Quick Hold)": "CAMPEONATO SUL-AMERICANO MASCULINO DE FUTSAL SUB-17/2026/NA",
                "Local": "Estádio Maracanã - Jornalista Mário Filho",
                "Local de Locução": "Narração By JB",
                "Elenco": "",
                "Narrador": "André Felipe",
                "Comentarista": "Vander Carioca",
                "Repórter": "",
                "Coordenador": "Tomaz Leão",
                "Produtor": "-",
            }])
            GeradorEscalasApp.gerar_html(object(), "André Felipe", df, tmp)
            html_path = next(Path(tmp).glob("escala_*.html"))
            html = html_path.read_text(encoding="utf-8")
            self.assertIn("Vander Carioca", html)
            self.assertNotIn("André Felipe ; Vander Carioca", html)
            self.assertIn("Oi André Felipe", html)
            self.assertIn("Escala consolidada 10/08/2026 a 10/08/2026", html)
            self.assertIn("Dúvidas ou problemas? É só nos procurar:", html)
            self.assertIn("Claudio Rolim", html)
            self.assertIn("(21) 99767-9446", html)
            self.assertIn("Leticia Alvares", html)
            self.assertIn("(21) 99645-5219", html)
            self.assertIn("Juliana Vasconcellos", html)
            self.assertIn("(21) 99027-9306", html)
            self.assertIn("Carlla Amara", html)
            self.assertIn("(21) 99242-1837", html)
            self.assertIn("Evelyn Zygiel", html)
            self.assertIn("(21) 97385-6772", html)
            self.assertIn("Luan Sanchez", html)
            self.assertIn("(21) 97521-2048", html)
            self.assertIn("Julia Silva Pereira (logística)", html)
            self.assertIn("(21) 97480-7758", html)
            self.assertIn("CAMPEONATO SUL-AMERICANO MASCULINO DE FUTSAL SUB-17/2026/NA", html)
            self.assertIn("EQUADOR X BRASIL", html)
            self.assertIn("Narração By JB", html)

    def test_transmissive_selection_row_wins_over_unmarked_continuation(self):
        row = pd.Series({
            "Plataforma": "Sportv",
            "Evento/Programa": "Seleção Sportv",
            "Produto (WO/Quick Hold)": "SELEÇÃO SPORTV/NA/NA",
            "Event Group": "",
            "Data_raw": "07/09/2026",
            "Início": "13:00",
            "Air Start Time": "13:00",
        })
        grades = pd.DataFrame([
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-09-07"), "Início": "12:55", "Pré": None, "Fim": "15:15", "Evento": "SELEÇÃO SPORTV - ALTERAÇÃO DURAÇÃO", "V/I": "V"},
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-09-07"), "Início": "13:30", "Pré": None, "Fim": "15:15", "Evento": "SELEÇÃO SPORTV", "V/I": None},
        ])
        match = _find_best_grade_match(row, grades)
        self.assertIsNotNone(match)
        self.assertEqual(str(match["Início"]), "12:55")
        self.assertEqual(str(match["V/I"]), "V")

    def test_sportv5_unnamed_block_is_flattened(self):
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'sportv5.xlsx'
            columns = [None] * 32
            columns[0] = 'DATA GRADE'
            columns[4] = 'DATA REAL'
            columns[7] = 'Data'
            columns[8] = 'Hora'
            columns[9] = 'V/I'
            columns[10] = 'Evento/Programa'
            columns[15] = 'Evento/Programa.1'
            columns[20] = 'Evento/Programa.2'
            columns[25] = 'Evento/Programa.3'
            channel_row = [None] * 32
            channel_row[4] = '2026-09-17'
            channel_row[28] = 'SPORTV5'
            event_row = [None, None, None, None, '2026-09-17', None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, '12:00', 'V', 'CIRCUITO MUNDIAL DE SURFE', 'ETAPA DE TRESTLES']
            rows = [
                [None] * 32,
                [None] * 32,
                columns,
                channel_row,
                event_row,
            ]
            pd.DataFrame(rows).to_excel(path, index=False, header=False)
            grades = flatten_sportv_grade(str(path))
            sportv5 = grades[grades['Plataforma'].astype(str).str.upper().eq('SPORTV5')]
            self.assertGreaterEqual(len(sportv5), 1)
            self.assertTrue(sportv5['Evento'].astype(str).str.contains('CIRCUITO MUNDIAL DE SURFE').any())

    def test_combate_keeps_a_confirmar_time_marker(self):
        from pathlib import Path
        from engine_grades import process_combate_grade
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'combate_confirmar.xlsx'
            pd.DataFrame([
                {'DATA': '2026-09-17', 'INÍCIO': 'A CONFIRMAR', 'FIM': 'A CONFIRMAR', 'PRÉ': 'X', 'EVENTO': 'POWER SLAP 23', 'COMBATE': 'VIVO', 'SPORTV': ''},
            ]).to_excel(path, index=False)
            grades = process_combate_grade(str(path))
            row = grades.iloc[0]
            self.assertEqual(str(row['Início']), 'A CONFIRMAR')
            self.assertEqual(str(row['Fim']), 'A CONFIRMAR')
            self.assertTrue(_is_time_to_confirmar(row['Início']))
            self.assertFalse(_is_time_to_confirmar('15:00'))

    def test_tv_globo_event_without_globo_grade_is_not_matched_to_sportv(self):
        row = pd.Series({
            "Plataforma": "",
            "Cliente": "TV GLOBO ED. RJ - CPC JORNALISMO",
            "Evento/Programa": "BDRJ - EXIBIÇÃO",
            "Produto (WO/Quick Hold)": "BOM DIA RIO",
            "Event Group": "",
            "Data_raw": "07/09/2026",
            "Início": "06:00",
            "Air Start Time": "06:00",
        })
        grades = pd.DataFrame([
            {"Plataforma": "SPORTV3", "Data": pd.Timestamp("2026-09-07"), "Início": "06:30", "Pré": None, "Fim": None, "Evento": "DIAMOND LEAGUE - 15ª ETAPA - DIA 1", "V/I": None},
        ])
        self.assertIsNone(_find_best_grade_match(row, grades))

    def test_generic_dia_does_not_create_event_match(self):
        grade = pd.Series({"Evento": "DIAMOND LEAGUE - 15ª ETAPA - DIA 1", "Início": "06:30", "Pré": None})
        self.assertEqual(_score_grade_match(grade, "BDRJ EXIBIÇÃO BOM DIA RIO", 6 * 3600), -100)

    def test_gecom_platform_matches_main_sportv_grade(self):
        row = pd.Series({
            "Plataforma": "GE.com 01",
            "Evento/Programa": "PANELA SPORTV",
            "Produto (WO/Quick Hold)": "PANELA SPORTV/NA/NA",
            "Event Group": "PANELA SPORTV",
            "Data_raw": "10/08/2026",
            "Início": "17:00",
            "Air Start Time": "17:00",
            "Tipo de Produção": "Exibição - Internet",
        })
        grades = pd.DataFrame([
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-08-10"), "Início": "23:30", "Pré": None, "Fim": "01:30", "Evento": "PANELA", "V/I": "I"},
        ])
        match = _find_best_grade_match(row, grades)
        self.assertIsNotNone(match)
        self.assertEqual(str(match["Plataforma"]), "SPORTV")
        self.assertEqual(str(match["Evento"]), "PANELA")

    def test_ppv_pre_hora_different_channel_attaches_by_matchup(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "ppv_separate_channel.xlsx")
            raw = pd.DataFrame([
                {"DATA": "2026-09-12", "HORA": "17:30", "CANAL": "PRE/GE TV", "EVENTO": "PRÉ-HORA", "MANDANTE": "PALMEIRAS", "VISITANTE": "SÃO PAULO", "PÓS": None},
                {"DATA": "2026-09-12", "HORA": "18:30", "CANAL": "PREMIERE", "EVENTO": "BRASILEIRO", "MANDANTE": "PALMEIRAS", "VISITANTE": "SÃO PAULO", "PÓS": "20:40"},
            ])
            raw.to_excel(path, index=False)
            result = process_premiere_grade(path)
            self.assertEqual(len(result), 1)
            self.assertEqual(_time_key(result.iloc[0]["Pré"]), 17 * 60 + 30)
            self.assertEqual(_time_key(result.iloc[0]["Início"]), 18 * 60 + 30)
            self.assertEqual(_time_key(result.iloc[0]["Fim"]), 20 * 60 + 40)

    def test_ppv_separate_pre_hora_attaches_to_following_event(self):
        from pathlib import Path
        ppv = Path('/home/ubuntu/escalas_regular/agosto_2026/PPV2026(Ago11ªversão).xlsx')
        if not ppv.exists():
            self.skipTest('Arquivo PPV de agosto não disponível')
        grades = process_premiere_grade(str(ppv))
        match = grades[(grades['Evento'].astype(str).str.upper().str.contains('FLUMINENSE X PALMEIRAS', na=False)) & (grades['Data'].astype(str).str.startswith('2026-08-15'))]
        self.assertEqual(len(match), 1)
        row = match.iloc[0]
        self.assertEqual(str(row['Início']), '16:30')
        self.assertEqual(str(row['Pré']), '15:30')
        self.assertEqual(str(row['Fim']), '18:40')

    def test_vt_inedito_futsal_matches_sportv2_same_day(self):
        row = pd.Series({
            "Plataforma": "Sportv 2",
            "Evento/Programa": "EQUADOR X BRASIL",
            "Produto (WO/Quick Hold)": "CAMPEONATO SUL-AMERICANO MASCULINO DE FUTSAL SUB-17/2026/NA",
            "Event Group": "CAMPEONATO SUL-AMERICANO MASCULINO DE FUTSAL SUB-17",
            "Data_raw": "10/08/2026",
            "Início": "12:00",
            "Air Start Time": "12:00",
            "Tipo de Produção": "Exibição - VT no Controle",
        })
        grades = pd.DataFrame([
            {"Plataforma": "SPORTV2", "Data": pd.Timestamp("2026-08-10"), "Início": "12:00", "Pré": None, "Fim": "13:45", "Evento": "CAMPEONATO SUL-AMERICANO MASCULINO DE FUTSAL SUB-17 - EQUADOR X BRASIL", "V/I": "I"},
        ])
        match = _find_best_grade_match(row, grades)
        self.assertIsNotNone(match)
        self.assertEqual(str(match["V/I"]), "I")
        self.assertEqual(str(match["Plataforma"]), "SPORTV2")

    def test_date_only_row_uses_report_start_time_and_same_day_match(self):
        row = pd.Series({
            "Plataforma": "Sportv",
            "Evento/Programa": "Sportv News",
            "Produto (WO/Quick Hold)": "SPORTV NEWS/1ª ED",
            "Event Group": "SPORTV NEWS",
            "Data_raw": "15/08/2026",
            "Início": "13:30",
            "Air Start Time": "13:30",
        })
        grades = pd.DataFrame([
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-08-14"), "Início": "21:30", "Pré": None, "Fim": "23:30", "Evento": "SPORTV NEWS", "V/I": "V"},
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-08-15"), "Início": "13:30", "Pré": None, "Fim": "14:00", "Evento": "SPORTV NEWS", "V/I": "V"},
        ])
        match = _find_best_grade_match(row, grades)
        self.assertIsNotNone(match)
        self.assertEqual(pd.to_datetime(match["Data"]).strftime("%Y-%m-%d"), "2026-08-15")
        self.assertEqual(str(match["Início"]), "13:30")

    def test_brazilian_report_date_matches_same_grade_date(self):
        row = pd.Series({
            "Plataforma": "Sportv",
            "Evento/Programa": "TROCA DE PASSES",
            "Produto (WO/Quick Hold)": "TROCA DE PASSES/NA/NA",
            "Event Group": "TROCA DE PASSES",
            "Data_raw": "10/08/2026",
        })
        grades = pd.DataFrame([
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-08-10"), "Início": "22:00", "Pré": None, "Fim": "23:30", "Evento": "TROCA DE PASSES", "V/I": "V"},
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-10-08"), "Início": "22:00", "Pré": None, "Fim": "23:30", "Evento": "TROCA DE PASSES", "V/I": "V"},
        ])
        match = _find_best_grade_match(row, grades)
        self.assertIsNotNone(match)
        self.assertEqual(pd.to_datetime(match["Data"]).strftime("%Y-%m-%d"), "2026-08-10")

    def test_combined_wo_bridges_two_adjacent_grade_blocks(self):
        # A WO do 2468 cobre as duas finais (17h-20h40) como uma atividade só,
        # mas a grade lista cada final em um marcador distinto e adjacente
        # (masculina 17h-19h, feminina 19h-20h40). O match deve estender o Fim
        # até o fim do segundo bloco em vez de cortar no fim do primeiro.
        row = pd.Series({
            "Plataforma": "Sportv 2",
            "Evento/Programa": "COPA DO MUNDO DE ESCALADA - FINAL FEMININA E MASCULINA BOULDER",
            "Produto (WO/Quick Hold)": "COPA DO MUNDO DE ESCALADA/2026/NA",
            "Event Group": "COPA DO MUNDO DE ESCALADA",
            "Data_raw": "20/10/2026",
            "Início": "17:00",
            "Fim": "20:40",
            "Air Start Time": "17:00",
        })
        grades = pd.DataFrame([
            {"Plataforma": "SPORTV2", "Data": pd.Timestamp("2026-10-20"), "Início": "17:00", "Pré": None, "Fim": "19:00", "Evento": "COPA DO MUNDO DE ESCALADA - FINAL MASCULINA BOULDER", "V/I": "I"},
            {"Plataforma": "SPORTV2", "Data": pd.Timestamp("2026-10-20"), "Início": "19:00", "Pré": None, "Fim": "20:40", "Evento": "COPA DO MUNDO DE ESCALADA - FINAL FEMININA BOULDER", "V/I": "I"},
        ])
        match = _find_best_grade_match(row, grades)
        self.assertIsNotNone(match)
        self.assertEqual(str(match["Início"]), "17:00")
        self.assertEqual(str(match["Fim"]), "20:40")

    def test_bridging_does_not_cross_into_unrelated_next_block(self):
        # Sem relação textual com o bloco seguinte, não deve estender o Fim
        # além do bloco encontrado, mesmo que o 2468 declare um Fim mais tarde.
        row = pd.Series({
            "Plataforma": "Sportv",
            "Evento/Programa": "TROCA DE PASSES",
            "Produto (WO/Quick Hold)": "TROCA DE PASSES/NA/NA",
            "Event Group": "TROCA DE PASSES",
            "Data_raw": "20/10/2026",
            "Início": "23:30",
            "Fim": "02:00",
            "Air Start Time": "23:30",
        })
        grades = pd.DataFrame([
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-10-20"), "Início": "23:30", "Pré": None, "Fim": "01:00", "Evento": "TROCA DE PASSES", "V/I": "V"},
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-10-20"), "Início": "01:00", "Pré": None, "Fim": "02:00", "Evento": "CAMPEONATO BRASILEIRO - FLAMENGO X FLUMINENSE", "V/I": "R"},
        ])
        match = _find_best_grade_match(row, grades)
        self.assertIsNotNone(match)
        self.assertEqual(str(match["Fim"]), "01:00")

    def test_ge_tv_channel_is_excluded(self):
        self.assertTrue(_is_ge_tv_row(pd.Series({"Canal": "GE TV", "Plataforma": "GE TV"})))
        self.assertTrue(_is_ge_tv_row(pd.Series({"Canal": "GE-TV"})))
        self.assertFalse(_is_ge_tv_row(pd.Series({"Canal": "Sportv", "Plataforma": "Sportv 2"})))

    def test_daytime_event_does_not_match_same_title_on_other_date(self):
        row = pd.Series({
            "Plataforma": "Sportv 3",
            "Evento/Programa": "COPA DO MUNDO DE CANOAGEM VELOCIDADE E PARACANOAGEM - SPRINT,PARACANOAGEM E SPRINT 5KM",
            "Produto (WO/Quick Hold)": "COPA DO MUNDO DE CANOAGEM DE VELOCIDADE/2026/NA",
            "Event Group": "COPA DO MUNDO DE CANOAGEM VELOCIDADE E PARACANOAGEM",
            "Data_raw": pd.Timestamp("2026-07-15 14:00:00"),
        })
        grades = pd.DataFrame([
            {
                "Plataforma": "SPORTV3",
                "Data": pd.Timestamp("2026-07-11"),
                "Início": "17:00",
                "Pré": None,
                "Fim": "19:00",
                "Evento": "COPA DO MUNDO DE CANOAGEM VELOCIDADE E PARACANOAGEM - SPRINT E PARACANOAGEM",
                "V/I": "I",
            }
        ])
        self.assertIsNone(_find_best_grade_match(row, grades))

    def test_generic_vt_grade_does_not_match_specific_event(self):
        grade = pd.Series({"Evento": "VT DE EVENTO", "Início": "14:00", "Pré": None})
        score = _score_grade_match(grade, "COPA DO MUNDO DE CANOAGEM VELOCIDADE E PARACANOAGEM", 14 * 3600)
        self.assertEqual(score, -100)

    def test_event_score_prefers_exact_match(self):
        row = pd.Series({"W/O Description": "LIGA DAS NAÇÕES MASCULINA DE VÔLEI - IRÃ X ALEMANHA"})
        self.assertEqual(
            _event_score("LIGA DAS NAÇÕES MASCULINA DE VÔLEI - IRÃ X ALEMANHA", row),
            1.0,
        )

    def test_folga_and_quick_hold_rules(self):
        folga = pd.Series({"Tipo de Atividade": "Booking", "Descrição": "FOLGA"})
        quick_hold = pd.Series({"Tipo de Atividade": "QUICK HOLD", "Descrição": "PODCAST"})
        viagem = pd.Series({"Tipo de Atividade": "QUICK HOLD", "Descrição": "PODCAST VIAGEM"})
        self.assertTrue(_is_folga_row(folga))
        self.assertTrue(_is_quickhold_in_scale(quick_hold))
        self.assertFalse(_is_quickhold_in_scale(viagem))

    def test_quickhold_includes_other_activity_types_not_just_podcast_cabine(self):
        # Regressão: a lista branca antiga só aceitava Quick Hold com "PODCAST"
        # ou "CABINE DO JOGO" no texto, descartando silenciosamente categorias
        # legítimas (gravações, participações em outros programas, reuniões).
        gravacao = pd.Series({"Tipo de Atividade": "Quick Hold", "Row Display": "Gravação DOC Acesso Total Futebol Feminino"})
        participacao = pd.Series({"Tipo de Atividade": "Quick Hold", "Row Display": "PARTICIPAÇÃO - MAIS VOCÊ"})
        reuniao = pd.Series({"Tipo de Atividade": "Quick Hold", "Row Display": "Papo remoto c/ Flavio Orro (GP Brasil F1)"})
        self.assertTrue(_is_quickhold_in_scale(gravacao))
        self.assertTrue(_is_quickhold_in_scale(participacao))
        self.assertTrue(_is_quickhold_in_scale(reuniao))

    def test_quickhold_2468_module_matches_cross_module_rule(self):
        # engine_2468._is_quickhold_in_scale decide quais Quick Hold sem WO#
        # sobrevivem ao dropna antes mesmo de chegar no filtro de engine_cross.
        # As duas cópias precisam ficar em sincronia; já houve regressão por
        # uma ter sido corrigida e a outra não.
        gravacao = pd.Series({"Tipo de Atividade": "Quick Hold", "Row Display": "Gravação DOC Acesso Total Futebol Feminino"})
        pessoal = pd.Series({"Tipo de Atividade": "Quick Hold", "Row Display": "COMPROMISSO PESSOAL"})
        viagem = pd.Series({"Tipo de Atividade": "Quick Hold", "Row Display": "VIAGEM IDA SP X RJ"})
        self.assertTrue(_is_quickhold_in_scale_2468(gravacao))
        self.assertFalse(_is_quickhold_in_scale_2468(pessoal))
        self.assertFalse(_is_quickhold_in_scale_2468(viagem))

    def test_quickhold_excludes_personal_and_internal_reservation(self):
        pessoal = pd.Series({"Tipo de Atividade": "Quick Hold", "Row Display": "COMPROMISSO PESSOAL"})
        evento_pessoal = pd.Series({"Tipo de Atividade": "Quick Hold", "Row Display": "Evento Pessoal"})
        reservado = pd.Series({"Tipo de Atividade": "Quick Hold", "Row Display": "Reservado Campanha Superliga 26/27"})
        folga_qh = pd.Series({"Tipo de Atividade": "Quick Hold", "Row Display": "folga de aniversario"})
        self.assertFalse(_is_quickhold_in_scale(pessoal))
        self.assertFalse(_is_quickhold_in_scale(evento_pessoal))
        self.assertFalse(_is_quickhold_in_scale(reservado))
        self.assertFalse(_is_quickhold_in_scale(folga_qh))

    def test_time_validation_rejects_undefined_values(self):
        self.assertTrue(_is_valid_time_str("01:30"))
        self.assertFalse(_is_valid_time_str("A Definir"))
        self.assertFalse(_is_valid_time_str("-"))

    def test_vt_in_grade_is_not_treated_as_program(self):
        row = pd.Series({
            "Tipo de Produção": "Exibição - VT no Controle",
            "Tipo de Atividade": "Booking",
            "Evento/Programa": "COPA DO MUNDO DE CANOAGEM VELOCIDADE E PARACANOAGEM",
        })
        grade = pd.Series({"Evento": "COPA DO MUNDO DE CANOAGEM VELOCIDADE E PARACANOAGEM", "V/I": "I"})
        self.assertFalse(_is_programa(row, grade))

    def test_vt_in_grade_with_live_indicator_is_not_treated_as_program(self):
        row = pd.Series({
            "Tipo de Produção": "Exibição - VT no Controle",
            "Tipo de Atividade": "Booking",
            "Evento/Programa": "EVENTO ESPORTIVO AO VIVO",
        })
        grade = pd.Series({"Evento": "EVENTO ESPORTIVO AO VIVO", "V/I": "V"})
        self.assertFalse(_is_programa(row, grade))

    def test_sportv_parser_uses_prejogo_as_previous_window_boundary(self):
        df = pd.DataFrame([
            ["SÁB", "2026-08-15", "13:30", "V", "SPORTV NEWS", "", "PROGRAMA", "SPORTV", "", 1/24, ""],
            ["SÁB", "2026-08-15", "14:00", "V", "PRÉ-JOGO", "", "PROGRAMA", "SPORTV", "", 1/48, ""],
        ], columns=["dia", "data", "hora", "vi", "evento", "obs", "tipo", "canal", "extra", "duracao", "extra2"])
        events = []
        extract_sportv_channel_block(df, 4, events)
        news = [event for event in events if event["Evento"] == "SPORTV NEWS"][0]
        self.assertEqual(str(news["Fim"]), "14:00")

    def test_consecutive_aquecimento_keeps_first_pre_time(self):
        df = pd.DataFrame([
            ["SÁB", "2026-09-12", "20:00", "V", "AQUECIMENTO SPORTV", "", "PROGRAMA", "SPORTV", "", 1/24, ""],
            ["SÁB", "2026-09-12", "20:30", None, "AQUECIMENTO SPORTV", "", "PROGRAMA", "SPORTV", "", 1/48, ""],
            ["SÁB", "2026-09-12", "21:00", "V", "CAMPEONATO BRASILEIRO DE FUTEBOL", "SANTOS X CRUZEIRO", "EVENTO", "SPORTV", "", 2/24, ""],
        ], columns=["dia", "data", "hora", "vi", "evento", "obs", "tipo", "canal", "extra", "duracao", "extra2"])
        events = []
        extract_sportv_channel_block(df, 4, events)
        match = [event for event in events if "SANTOS X CRUZEIRO" in event["Evento"]][0]
        self.assertEqual(str(match["Pré"]), "20:00")
        self.assertEqual(str(match["Início"]), "21:00")

    def test_repeated_sportv_news_window_reaches_1430(self):
        grades = pd.DataFrame([
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-08-15"), "Início": "13:30", "Fim": "14:30", "Evento": "SPORTV NEWS", "V/I": "V"},
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-08-15"), "Início": "14:00", "Fim": "14:30", "Evento": "SPORTV NEWS", "V/I": "I"},
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-08-15"), "Início": "14:30", "Fim": "14:30", "Evento": "PRÉ-JOGO", "V/I": "V"},
        ])
        consolidated = _consolidate_grade_dataframe_windows(grades)
        news = consolidated[consolidated["Evento"] == "SPORTV NEWS"].iloc[0]
        self.assertEqual(str(news["Início"]), "13:30")
        self.assertEqual(str(news["Fim"]), "14:30")

    def test_reprise_distinct_event_closes_previous_window(self):
        events = [
            {"Plataforma": "SPORTV2", "Data": "2026-09-12", "Início": "10:30", "Pré": "10:10", "Fim": "13:00", "Evento": "PRÉ-OLÍMPICO FEMININO DE VÔLEI - ARGENTINA X VENEZUELA", "V/I": "V"},
            {"Plataforma": "SPORTV2", "Data": "2026-09-12", "Início": "11:00", "Pré": None, "Fim": "12:30", "Evento": "PRÉ-OLÍMPICO FEMININO DE VÔLEI - PRIMEIRA FASE", "V/I": None},
            {"Plataforma": "SPORTV2", "Data": "2026-09-12", "Início": "12:30", "Pré": None, "Fim": "13:00", "Evento": "SESSÃO SPORTV - FOME DE MEDALHA - VÔLEI", "V/I": "R"},
            {"Plataforma": "SPORTV2", "Data": "2026-09-12", "Início": "13:30", "Pré": "13:00", "Fim": "16:10", "Evento": "PRÉ-OLÍMPICO FEMININO DE VÔLEI - BRASIL X CHILE", "V/I": "V"},
        ]
        extended = _extend_grade_windows_to_next_event(events)
        self.assertEqual(extended[0]["Fim"], "12:30")

    def test_dataframe_window_uses_next_event_boundary(self):
        grades = pd.DataFrame([
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-08-10"), "Início": "22:00", "Fim": "22:30", "Evento": "TROCA DE PASSES", "V/I": "V"},
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-08-10"), "Início": "22:30", "Fim": "23:30", "Evento": "HELLO LA", "V/I": "R"},
            {"Plataforma": "SPORTV", "Data": pd.Timestamp("2026-08-10"), "Início": "23:30", "Fim": "23:30", "Evento": "PANELA", "V/I": "I"},
        ])
        consolidated = _consolidate_grade_dataframe_windows(grades)
        troca = consolidated[consolidated["Evento"] == "TROCA DE PASSES"].iloc[0]
        self.assertEqual(str(troca["Fim"]), "22:30")

    def test_merge_does_not_swallow_post_midnight_rerun_of_earlier_title(self):
        # Mesmo titulo as 19:00 (ao vivo) e novamente as 01:00 (reprise, ja
        # virada a meia-noite) no mesmo dia de transmissao. Comparar horarios
        # brutos faria 01:00 parecer "dentro" da janela das 19:00-21:00 (pois
        # 01:00 < 21:00 em minutos), descartando a reprise por engano e
        # estendendo o TROCA DE PASSES das 23:30 ate o evento seguinte (02:00)
        # em vez de parar no evento real das 01:00.
        events = [
            {"Plataforma": "SPORTV", "Data": "2026-10-15", "Início": "19:00", "Fim": "21:00", "Evento": "COPA - FLAMENGO X FLUMINENSE", "V/I": "V"},
            {"Plataforma": "SPORTV", "Data": "2026-10-15", "Início": "23:00", "Fim": "23:30", "Evento": "GIRO DA RODADA", "V/I": "R"},
            {"Plataforma": "SPORTV", "Data": "2026-10-15", "Início": "23:30", "Fim": "01:00", "Evento": "TROCA DE PASSES", "V/I": "V"},
            {"Plataforma": "SPORTV", "Data": "2026-10-15", "Início": "00:00", "Fim": "01:00", "Evento": "TROCA DE PASSES", "V/I": None},
            {"Plataforma": "SPORTV", "Data": "2026-10-15", "Início": "00:30", "Fim": "01:00", "Evento": "TROCA DE PASSES", "V/I": None},
            {"Plataforma": "SPORTV", "Data": "2026-10-15", "Início": "01:00", "Fim": "02:00", "Evento": "COPA - FLAMENGO X FLUMINENSE", "V/I": "R"},
            {"Plataforma": "SPORTV", "Data": "2026-10-15", "Início": "02:00", "Fim": "04:00", "Evento": "COPA SUB-20 - OITAVAS", "V/I": "R"},
        ]
        merged = _merge_repeated_grade_windows(events)
        copa_entries = [e for e in merged if e["Evento"] == "COPA - FLAMENGO X FLUMINENSE"]
        self.assertEqual(len(copa_entries), 2)
        rerun = [e for e in copa_entries if e["Início"] == "01:00"][0]
        self.assertEqual(rerun["Fim"], "02:00")

        extended = _extend_grade_windows_to_next_event(merged)
        troca = [e for e in extended if e["Evento"] == "TROCA DE PASSES"][0]
        self.assertEqual(troca["Fim"], "01:00")

    def test_program_window_extends_to_next_distinct_event(self):
        events = [
            {"Plataforma": "SPORTV", "Data": "2026-08-10", "Início": "22:00", "Fim": "22:30", "Evento": "TROCA DE PASSES", "V/I": "V"},
            {"Plataforma": "SPORTV", "Data": "2026-08-10", "Início": "23:30", "Fim": "23:30", "Evento": "PANELA", "V/I": "I"},
        ]
        extended = _extend_grade_windows_to_next_event(events)
        self.assertEqual(extended[0]["Fim"], "23:30")

    def test_repeated_uts_markers_are_merged_until_next_distinct_event(self):
        events = [
            {"Plataforma": "SPORTV3", "Data": "2026-07-16", "Início": "18:00", "Fim": "19:00", "Evento": "UTS - 1ª RODADA", "V/I": "V"},
            {"Plataforma": "SPORTV3", "Data": "2026-07-16", "Início": "18:30", "Fim": "19:00", "Evento": "UTS - RIO DE JANEIRO, BRASIL", "V/I": None},
            {"Plataforma": "SPORTV3", "Data": "2026-07-16", "Início": "19:00", "Fim": "20:00", "Evento": "UTS - 1ª RODADA", "V/I": "V"},
            {"Plataforma": "SPORTV3", "Data": "2026-07-16", "Início": "19:30", "Fim": "20:00", "Evento": "UTS - RIO DE JANEIRO, BRASIL", "V/I": None},
            {"Plataforma": "SPORTV3", "Data": "2026-07-16", "Início": "20:00", "Fim": "21:00", "Evento": "UTS - 1ª RODADA", "V/I": "V"},
            {"Plataforma": "SPORTV3", "Data": "2026-07-16", "Início": "20:30", "Fim": "21:00", "Evento": "UTS - RIO DE JANEIRO, BRASIL", "V/I": None},
            {"Plataforma": "SPORTV3", "Data": "2026-07-16", "Início": "21:00", "Fim": "22:00", "Evento": "UTS - 1ª RODADA", "V/I": "V"},
            {"Plataforma": "SPORTV3", "Data": "2026-07-16", "Início": "21:30", "Fim": "22:00", "Evento": "UTS - RIO DE JANEIRO, BRASIL", "V/I": None},
            {"Plataforma": "SPORTV3", "Data": "2026-07-16", "Início": "22:00", "Fim": "23:00", "Evento": "LIGA DAS NAÇÕES MASCULINA DE VÔLEI - EUA X BRASIL", "V/I": "V"},
        ]
        merged = _merge_repeated_grade_windows(events)
        uts = [event for event in merged if event["Evento"] == "UTS - 1ª RODADA"]
        self.assertEqual(len(uts), 1)
        self.assertEqual(uts[0]["Início"], "18:00")
        self.assertEqual(uts[0]["Fim"], "22:00")
        self.assertEqual(len(merged), 6)

    def test_vt_grade_without_vi_indicator_is_treated_as_program(self):
        row = pd.Series({
            "Tipo de Produção": "Exibição - VT no Controle",
            "Tipo de Atividade": "Booking",
            "Evento/Programa": "VT DE EVENTO ESPECIAL",
        })
        grade = pd.Series({"Evento": "VT DE EVENTO ESPECIAL", "V/I": None})
        self.assertTrue(_is_programa(row, grade))

    def test_vt_without_grade_remains_program(self):
        row = pd.Series({
            "Tipo de Produção": "Exibição - VT no Controle",
            "Tipo de Atividade": "Booking",
            "Evento/Programa": "VT DE EVENTO ESPECIAL",
        })
        self.assertTrue(_is_programa(row, None))

    def test_event_without_pre_requires_review(self):
        out_row = {"Pré": "-", "Início": "22:00", "Fim": "02:00"}
        alerts = []
        severity = _append_pre_review_alerts(out_row, alerts, "OK", is_prog=False, has_valid_grade_time=True)
        self.assertEqual(severity, "YELLOW")
        self.assertEqual(alerts, ["Conferir Pré"])

    def test_missing_grade_is_flagged_even_for_program(self):
        out_row = {"Pré": "", "Início": "09:00", "Fim": "13:00"}
        alerts = []
        severity = _mark_missing_grade(out_row, alerts, "OK")
        self.assertEqual(severity, "YELLOW")
        self.assertEqual(alerts, ["Horário não encontrado na Grade"])
        self.assertEqual(out_row["Início"], "09:00")
        self.assertEqual(out_row["Fim"], "13:00")
        self.assertEqual(out_row["Pré"], "-")

    def test_program_without_pre_remains_ok(self):
        out_row = {"Pré": "-", "Início": "22:00", "Fim": "02:00"}
        alerts = []
        severity = _append_pre_review_alerts(out_row, alerts, "OK", is_prog=True, has_valid_grade_time=True)
        self.assertEqual(severity, "OK")
        self.assertEqual(alerts, [])

    def test_event_with_pre_does_not_require_pre_review(self):
        out_row = {"Pré": "21:30", "Início": "22:00", "Fim": "02:00"}
        alerts = []
        severity = _append_pre_review_alerts(out_row, alerts, "OK", is_prog=False, has_valid_grade_time=True)
        self.assertEqual(severity, "OK")
        self.assertEqual(alerts, [])

    def test_one_fight_night_45_matches_combate_grade(self):
        base_row = pd.Series(
            {
                "Plataforma": "Combate",
                "Data_raw": pd.Timestamp("2026-07-17 22:00:00"),
                "Evento/Programa": "ONE FIGHT NIGHT 45",
                "Produto (WO/Quick Hold)": "ONE FRIDAY FIGHTS/2026/NA",
                "Event Group": "ONE FIGHT NIGHT",
            }
        )
        grades = pd.DataFrame(
            [
                {
                    "Plataforma": "COMBATE",
                    "Data": pd.Timestamp("2026-07-17"),
                    "Início": "22:00:00",
                    "Pré": "X",
                    "Fim": "02:00:00",
                    "Evento": "ONE FIGHT NIGHT 45",
                    "V/I": "V",
                }
            ]
        )
        match = _find_best_grade_match(base_row, grades)
        self.assertIsNotNone(match)
        self.assertEqual(match["Evento"], "ONE FIGHT NIGHT 45")

    def test_autofill_recognizes_contacts_and_check_files(self):
        from gerador_escalas_desktop import GeradorEscalasApp
        class DummyEntry:
            def __init__(self, val=""):
                self.val = val
            def get(self):
                return self.val
            def delete(self, s, e):
                self.val = ""
            def insert(self, idx, val):
                self.val = val

        with tempfile.TemporaryDirectory() as tmpdir:
            contacts_file = Path(tmpdir) / "contatos_nova_versao.xlsx"
            check_file = Path(tmpdir) / "Check_Pre_Envio_Gerado.xlsx"
            contacts_file.write_bytes(b"dummy")
            check_file.write_bytes(b"dummy")

            app = object.__new__(GeradorEscalasApp)
            app.base_dir = tmpdir
            app.config = {}
            app.entry_2405 = DummyEntry()
            app.entry_2468 = DummyEntry()
            app.entry_check = DummyEntry()
            app.entry_contacts = DummyEntry()

            GeradorEscalasApp._autofill_default_paths(app)

            self.assertEqual(app.entry_contacts.get(), str(contacts_file))
            self.assertEqual(app.entry_check.get(), str(check_file))

    def test_sort_dataframe_by_date_and_time_chronological_order(self):
        from gerador_escalas_desktop import _sort_dataframe_by_date_and_time
        df_unordered = pd.DataFrame([
            {"Nome": "PC Oliveira", "Data": "21/09/2026", "Início": "20:30", "Evento/Descrição": "CORINTHIANS X BAHIA"},
            {"Nome": "PC Oliveira", "Data": "25/09/2026", "Início": "06:40", "Evento/Descrição": "BRASIL X AUSTRÁLIA"},
            {"Nome": "PC Oliveira", "Data": "22/09/2026", "Início": "-", "Evento/Descrição": "FOLGA"},
            {"Nome": "PC Oliveira", "Data": "23/09/2026", "Início": "-", "Evento/Descrição": "FOLGA"},
        ])
        df_sorted = _sort_dataframe_by_date_and_time(df_unordered)
        dates_result = df_sorted["Data"].tolist()
        self.assertEqual(dates_result, ["21/09/2026", "22/09/2026", "23/09/2026", "25/09/2026"])

    def test_html_period_uses_the_selected_filter_not_the_last_event(self):
        """O período divulgado é a janela escolhida na Etapa 3.

        Quando a escala do profissional termina antes do fim do filtro, o
        período do HTML (e por consequência o assunto do e-mail) precisa
        continuar mostrando a data final selecionada pelo usuário.
        """
        from gerador_escalas_desktop import (
            _date_range_bounds,
            _format_period_dates,
            _split_scale_weeks,
        )

        def build(dates):
            df = pd.DataFrame([{"Nome": "PC Oliveira", "Data": d, "Início": "19:00"} for d in dates])
            bounds = df["Data"].apply(_date_range_bounds)
            df["Data_obj"] = bounds.apply(lambda item: item[0])
            df["Data_fim_obj"] = bounds.apply(lambda item: item[1])
            return df

        filtro_inicio = pd.to_datetime("28/09/2026", dayfirst=True)
        filtro_fim = pd.to_datetime("04/10/2026", dayfirst=True)

        # Último compromisso em 03/10, mas o filtro vai até 04/10.
        df = build(["28/09/2026", "01/10/2026", "02/10/2026", "03/10/2026"])
        semanas = _split_scale_weeks(df, filtro_inicio, filtro_fim)
        self.assertEqual(len(semanas), 1)
        _, _, period_start, period_end = semanas[0]
        self.assertEqual(_format_period_dates(period_start, period_end), "28/09/2026 a 04/10/2026")

        # Sem filtro informado, o comportamento antigo (bordas dos dados) é mantido.
        _, _, sem_filtro_inicio, sem_filtro_fim = _split_scale_weeks(df)[0]
        self.assertEqual(
            _format_period_dates(sem_filtro_inicio, sem_filtro_fim),
            "28/09/2026 a 03/10/2026",
        )

        # Janela de duas semanas: os blocos seguem o filtro e nenhuma linha some,
        # nem a que começa antes do início do filtro.
        df_longo = build(["26/09/2026 a 29/09/2026", "02/10/2026", "10/10/2026"])
        semanas_longo = _split_scale_weeks(
            df_longo, filtro_inicio, pd.to_datetime("11/10/2026", dayfirst=True)
        )
        rotulos = [_format_period_dates(ini, fim) for _, _, ini, fim in semanas_longo]
        self.assertEqual(rotulos, ["28/09/2026 a 04/10/2026", "05/10/2026 a 11/10/2026"])
        self.assertEqual(sum(len(part) for _, part, _, _ in semanas_longo), len(df_longo))

    def test_status_revisao_color_follows_the_row_after_sorting(self):
        """A cor da célula de Status Revisão tem que acompanhar a própria linha.

        A ordenação final por Nome/Data/Início reordena as linhas depois que a
        severidade é calculada; se o mapa de cores for lido pela posição antiga,
        linhas "OK" saem amarelas e linhas com alerta saem verdes.
        """
        from openpyxl import load_workbook

        from engine_cross import run_etapa1

        # A linha OK entra primeiro, mas a ordenação por Nome joga a linha de
        # alerta para a frente, trocando as posições das duas.
        rows = [
            {"WO#": "", "Nome": "ZULMIRA TESTE", "Tipo de Atividade": "Folga",
             "Descrição": "Day Off / Folga", "Data": "12/09/2026", "Início": "08:00",
             "Função": "Narrador", "Local de Locução": "ION"},
            {"WO#": "", "Nome": "ANA TESTE", "Tipo de Atividade": "Quick Hold",
             "Descrição": "Podcast", "Data": "12/09/2026", "Início": "09:00",
             "Função": "Narrador", "Local de Locução": ""},
        ]

        with tempfile.TemporaryDirectory() as tmpdir:
            path_2468 = os.path.join(tmpdir, "2468.xlsx")
            pd.DataFrame(rows).to_excel(path_2468, index=False)

            out_path = run_etapa1(path_2468, None, None, None, None, None, None)

            wb = load_workbook(out_path)
            ws = wb.active
            header = [cell.value for cell in ws[1]]
            status_idx = header.index("Status Revisão") + 1
            name_idx = header.index("Nome") + 1

            painted = {
                ws.cell(row=r, column=name_idx).value: (
                    ws.cell(row=r, column=status_idx).value,
                    ws.cell(row=r, column=status_idx).fill.start_color.rgb,
                )
                for r in range(2, ws.max_row + 1)
            }

            self.assertEqual(painted["ANA TESTE"], ("Local Ausente", "FFFFFFCC"))
            self.assertEqual(painted["ZULMIRA TESTE"], ("OK", "FF00B050"))

            # A amostra da aba Legenda precisa usar a mesma cor da planilha.
            legend = wb["Legenda"]
            legend_colors = {
                legend.cell(row=r, column=1).value: legend.cell(row=r, column=1).fill.start_color.rgb
                for r in range(4, legend.max_row + 1)
            }
            self.assertEqual(legend_colors["OK"], "FF00B050")
            self.assertEqual(legend_colors["Local Ausente"], "FFFFFFCC")
            self.assertIn("Aquecimento Delas", legend_colors)
            self.assertEqual(legend_colors["Aquecimento Delas"], "FFE1D5E7")

    def test_aquecimento_delas_triggers_specific_review_status_and_legend(self):
        from openpyxl import load_workbook
        from engine_cross import run_etapa1

        # Grade com Aquecimento Delas seguido por jogo feminino
        df_grade = pd.DataFrame([
            ["QUI", "2026-10-15", "16:00", "V", "AQUECIMENTO DELAS", "GLOBOPLAY", "PROGRAMA", "SPORTV", "", 1/48, ""],
            ["QUI", "2026-10-15", "16:30", None, "AQUECIMENTO DELAS", "", "PROGRAMA", "SPORTV", "", 1/48, ""],
            ["QUI", "2026-10-15", "17:00", "V", "CONMEBOL LIBERTADORES FEMININA", "CORINTHIANS X COLO COLO", "EVENTO", "SPORTV", "", 2/24, ""],
            ["QUI", "2026-10-15", "19:00", "V", "COPA DO BRASIL FEMININA", "FLAMENGO X FLUMINENSE", "EVENTO", "SPORTV", "", 2/24, ""],
        ], columns=["dia", "data", "hora", "vi", "evento", "obs", "tipo", "canal", "extra", "duracao", "extra2"])

        events = []
        extract_sportv_channel_block(df_grade, 4, events)
        match_ev = [ev for ev in events if "CORINTHIANS X COLO COLO" in ev["Evento"]][0]
        self.assertEqual(str(match_ev["Pré"]), "16:00")
        self.assertEqual(str(match_ev["Início"]), "17:00")
        self.assertEqual(str(match_ev["Pré_Origem"]), "AQUECIMENTO DELAS - GLOBOPLAY")

        # Base 2468 com narradora escalada para o jogo e narrador para outro jogo com aquecimento normal
        rows_2468 = [
            {"WO#": "1111111-1", "Nome": "NARRADORA TESTE", "Data": "15/10/2026", "Início": "17:00", "Fim": "19:00",
             "Evento/Programa": "CONMEBOL LIBERTADORES FEMININA - CORINTHIANS X COLO COLO",
             "Canal (Master Room)": "SPORTV", "Tipo de Produção": "AO VIVO", "Tipo Atividade": "BOOKING",
             "Função": "Narrador", "Local de Locução": "CABINE 1"},
            {"WO#": "2222222-1", "Nome": "NARRADOR NORMAL", "Data": "12/09/2026", "Início": "21:00", "Fim": "23:00",
             "Evento/Programa": "CAMPEONATO BRASILEIRO - SANTOS X CRUZEIRO",
             "Canal (Master Room)": "SPORTV", "Tipo de Produção": "AO VIVO", "Tipo Atividade": "BOOKING",
             "Função": "Narrador", "Local de Locução": "CABINE 2"},
        ]

        df_sp_normal = pd.DataFrame([
            ["SÁB", "2026-09-12", "20:00", "V", "AQUECIMENTO SPORTV", "", "PROGRAMA", "SPORTV", "", 1/24, ""],
            ["SÁB", "2026-09-12", "20:30", None, "AQUECIMENTO SPORTV", "", "PROGRAMA", "SPORTV", "", 1/48, ""],
            ["SÁB", "2026-09-12", "21:00", "V", "CAMPEONATO BRASILEIRO", "SANTOS X CRUZEIRO", "EVENTO", "SPORTV", "", 2/24, ""],
            ["SÁB", "2026-09-12", "23:00", "V", "TROCA DE PASSES", "", "PROGRAMA", "SPORTV", "", 1/24, ""],
        ], columns=["dia", "data", "hora", "vi", "evento", "obs", "tipo", "canal", "extra", "duracao", "extra2"])

        with tempfile.TemporaryDirectory() as tmpdir:
            path_2468 = os.path.join(tmpdir, "2468.xlsx")
            pd.DataFrame(rows_2468).to_excel(path_2468, index=False)

            # Salvar grade Sportv contendo ambos os dias
            df_sp_combined = pd.concat([df_grade, df_sp_normal], ignore_index=True)
            empty_cols = pd.DataFrame([[""] * 10] * len(df_sp_combined))
            df_full_grade = pd.concat([empty_cols, df_sp_combined], axis=1)
            path_sp = os.path.join(tmpdir, "grade_sportv.xlsx")
            df_full_grade.to_excel(path_sp, index=False)

            out_path = run_etapa1(path_2468, path_sp, None, None, None, None, None)

            wb = load_workbook(out_path)
            ws = wb.active
            header = [cell.value for cell in ws[1]]
            status_idx = header.index("Status Revisão") + 1
            pre_idx = header.index("Pré") + 1
            name_idx = header.index("Nome") + 1

            rows_out = {}
            for r in range(2, ws.max_row + 1):
                name = ws.cell(row=r, column=name_idx).value
                rows_out[name] = {
                    "status": ws.cell(row=r, column=status_idx).value,
                    "status_color": ws.cell(row=r, column=status_idx).fill.start_color.rgb,
                    "pre": ws.cell(row=r, column=pre_idx).value,
                }

            # Caso AQUECIMENTO DELAS: Pré 16:00, Status "Aquecimento Delas", cor FFE1D5E7
            self.assertEqual(rows_out["NARRADORA TESTE"]["pre"], "16:00")
            self.assertEqual(rows_out["NARRADORA TESTE"]["status"], "Aquecimento Delas")
            self.assertEqual(rows_out["NARRADORA TESTE"]["status_color"], "FFE1D5E7")

            # Demais situações de aquecimento continuam iguais: Pré 20:00, Status OK, cor FF00B050
            self.assertEqual(rows_out["NARRADOR NORMAL"]["pre"], "20:00")
            self.assertEqual(rows_out["NARRADOR NORMAL"]["status"], "OK")
            self.assertEqual(rows_out["NARRADOR NORMAL"]["status_color"], "FF00B050")

            # Aba Legenda contém Aquecimento Delas com a mesma cor
            legend = wb["Legenda"]
            legend_dict = {
                legend.cell(row=r, column=1).value: (
                    legend.cell(row=r, column=2).value,
                    legend.cell(row=r, column=1).fill.start_color.rgb,
                )
                for r in range(4, legend.max_row + 1)
            }
            self.assertIn("Aquecimento Delas", legend_dict)
            self.assertEqual(legend_dict["Aquecimento Delas"][1], "FFE1D5E7")
            self.assertIn("Aquecimento Delas", legend_dict["Aquecimento Delas"][0])

    def test_tv_globo_triggers_specific_review_status_and_legend(self):
        from openpyxl import load_workbook
        from engine_cross import run_etapa1

        rows_2468 = [
            {"WO#": "3333333-1", "Nome": "NARRADOR GLOBO", "Data": "18/10/2026", "Início": "16:00", "Fim": "18:00",
             "Evento/Programa": "CAMPEONATO BRASILEIRO - FLAMENGO X PALMEIRAS",
             "Canal (Master Room)": "TV GLOBO", "Tipo de Produção": "AO VIVO", "Tipo Atividade": "BOOKING",
             "Função": "Narrador", "Local de Locução": "MARACANÃ"},
            {"WO#": "4444444-1", "Nome": "COMENTARISTA GLOBO CLIENTE", "Data": "19/10/2026", "Início": "06:00", "Fim": "07:30",
             "Evento/Programa": "BOM DIA RIO", "Cliente": "TV GLOBO ED. RJ",
             "Canal (Master Room)": "", "Tipo de Produção": "AO VIVO", "Tipo Atividade": "BOOKING",
             "Função": "Comentarista", "Local de Locução": "ESTÚDIO"},
        ]

        df_sp = pd.DataFrame([
            ["DOM", "2026-10-18", "14:00", "V", "TROCA DE PASSES", "", "PROGRAMA", "SPORTV", "", 1/24, ""],
        ], columns=["dia", "data", "hora", "vi", "evento", "obs", "tipo", "canal", "extra", "duracao", "extra2"])

        with tempfile.TemporaryDirectory() as tmpdir:
            path_2468 = os.path.join(tmpdir, "2468.xlsx")
            pd.DataFrame(rows_2468).to_excel(path_2468, index=False)

            empty_cols = pd.DataFrame([[""] * 10] * len(df_sp))
            df_full_grade = pd.concat([empty_cols, df_sp], axis=1)
            path_sp = os.path.join(tmpdir, "grade_sportv.xlsx")
            df_full_grade.to_excel(path_sp, index=False)

            out_path = run_etapa1(path_2468, path_sp, None, None, None, None, None)

            wb = load_workbook(out_path)
            ws = wb.active
            header = [cell.value for cell in ws[1]]
            status_idx = header.index("Status Revisão") + 1
            pre_idx = header.index("Pré") + 1
            name_idx = header.index("Nome") + 1

            rows_out = {}
            for r in range(2, ws.max_row + 1):
                name = ws.cell(row=r, column=name_idx).value
                rows_out[name] = {
                    "status": ws.cell(row=r, column=status_idx).value,
                    "status_color": ws.cell(row=r, column=status_idx).fill.start_color.rgb,
                    "pre": ws.cell(row=r, column=pre_idx).value,
                }

            # Caso Canal TV GLOBO: Pré vazio/"-", Status "TV Globo", cor FFB4C6E7
            self.assertIn(rows_out["NARRADOR GLOBO"]["pre"], {None, "-", ""})
            self.assertEqual(rows_out["NARRADOR GLOBO"]["status"], "TV Globo")
            self.assertEqual(rows_out["NARRADOR GLOBO"]["status_color"], "FFB4C6E7")

            # Caso Cliente TV GLOBO: Pré vazio/"-", Status "TV Globo", cor FFB4C6E7
            self.assertIn(rows_out["COMENTARISTA GLOBO CLIENTE"]["pre"], {None, "-", ""})
            self.assertEqual(rows_out["COMENTARISTA GLOBO CLIENTE"]["status"], "TV Globo")
            self.assertEqual(rows_out["COMENTARISTA GLOBO CLIENTE"]["status_color"], "FFB4C6E7")

            # Aba Legenda contém TV Globo com a mesma cor FFB4C6E7
            legend = wb["Legenda"]
            legend_dict = {
                legend.cell(row=r, column=1).value: (
                    legend.cell(row=r, column=2).value,
                    legend.cell(row=r, column=1).fill.start_color.rgb,
                )
                for r in range(4, legend.max_row + 1)
            }
            self.assertIn("TV Globo", legend_dict)
            self.assertEqual(legend_dict["TV Globo"][1], "FFB4C6E7")
            self.assertIn("TV Globo", legend_dict["TV Globo"][0])


if __name__ == "__main__":
    unittest.main()
